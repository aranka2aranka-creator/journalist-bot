# -*- coding: utf-8 -*-
"""
admin_panel/app.py
-------------------
پنل مدیریت وب برای دیدن خبرهای ارسالی، اصلاح امتیاز (از جمله تایید صفحه اول)،
و گرفتن خروجی اکسل. با نام کاربری/رمز ساده محافظت می‌شود (برای شروع کافی است).
"""
import io
import os
import sys
import datetime
import functools

import requests
from flask import (
    Flask,
    render_template,
    request,
    redirect,
    url_for,
    session as flask_session,
    send_file,
    flash,
    Response,
    abort,
)

# اجازه بده db.py و scoring.py از پوشه‌ی اصلی پروژه import بشن
sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from db import Submission, get_session, init_db  # noqa: E402
from scoring import FRONT_PAGE_BONUS_SUGGESTION  # noqa: E402

app = Flask(__name__)
app.secret_key = os.environ.get("ADMIN_SECRET_KEY", "change-this-secret-key")

ADMIN_USERNAME = os.environ.get("ADMIN_USERNAME", "admin")
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "")
BALE_BOT_TOKEN = os.environ.get("BALE_BOT_TOKEN", "")
BALE_API_BASE = "https://tapi.bale.ai"

init_db()


def login_required(view):
    @functools.wraps(view)
    def wrapped(*args, **kwargs):
        if not flask_session.get("logged_in"):
            return redirect(url_for("login", next=request.path))
        return view(*args, **kwargs)

    return wrapped


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = request.form.get("username", "")
        password = request.form.get("password", "")
        if username == ADMIN_USERNAME and ADMIN_PASSWORD and password == ADMIN_PASSWORD:
            flask_session["logged_in"] = True
            next_url = request.args.get("next") or url_for("index")
            return redirect(next_url)
        flash("نام کاربری یا رمز عبور اشتباه است.")
    return render_template("login.html")


@app.route("/logout")
def logout():
    flask_session.clear()
    return redirect(url_for("login"))


@app.route("/")
@login_required
def index():
    db = get_session()
    try:
        newspaper_filter = request.args.get("newspaper", "").strip()
        query = db.query(Submission).order_by(Submission.created_at.desc())
        if newspaper_filter:
            query = query.filter(Submission.newspaper_name == newspaper_filter)
        submissions = query.all()

        totals: dict[str, int] = {}
        for s in db.query(Submission).all():
            totals[s.newspaper_name] = totals.get(s.newspaper_name, 0) + s.effective_score
        totals_sorted = sorted(totals.items(), key=lambda x: x[1], reverse=True)

        newspapers = sorted({s.newspaper_name for s in db.query(Submission).all()})

        return render_template(
            "index.html",
            submissions=submissions,
            totals=totals_sorted,
            newspapers=newspapers,
            active_filter=newspaper_filter,
        )
    finally:
        db.close()


@app.route("/submission/<int:submission_id>/update", methods=["POST"])
@login_required
def update_submission(submission_id):
    db = get_session()
    try:
        submission = db.query(Submission).get(submission_id)
        if submission is None:
            flash("این ارسال پیدا نشد.")
            return redirect(url_for("index"))

        is_front_page = request.form.get("is_front_page")
        final_score = request.form.get("final_score", "").strip()
        admin_note = request.form.get("admin_note", "").strip()

        submission.is_front_page = True if is_front_page == "yes" else (
            False if is_front_page == "no" else None
        )
        submission.final_score = int(final_score) if final_score else None
        submission.admin_note = admin_note or None
        submission.reviewed_at = datetime.datetime.utcnow()

        db.commit()
        flash("امتیاز با موفقیت ذخیره شد.")
    finally:
        db.close()
    return redirect(url_for("index"))


@app.route("/submission/<int:submission_id>/file")
@login_required
def download_file(submission_id):
    """فایلی که خبرنگار توی بله فرستاده (PDF یا عکس) رو از سرور بله می‌گیره
    و برای دانلود/نمایش توی مرورگر تحویل می‌ده."""
    if not BALE_BOT_TOKEN:
        abort(500, "متغیر BALE_BOT_TOKEN برای این سرویس تنظیم نشده، پس نمی‌شود فایل را از بله گرفت.")

    db = get_session()
    try:
        submission = db.query(Submission).get(submission_id)
        if submission is None:
            abort(404)
        file_id = submission.file_id
        file_kind = submission.file_kind
        newspaper_name = submission.newspaper_name
    finally:
        db.close()

    info_resp = requests.get(
        f"{BALE_API_BASE}/bot{BALE_BOT_TOKEN}/getFile",
        params={"file_id": file_id},
        timeout=20,
    )
    info = info_resp.json()
    if not info.get("ok"):
        abort(502, "بله فایل را برنگرداند؛ ممکن است لینک منقضی شده یا فایل حذف شده باشد.")

    file_path = info["result"]["file_path"]

    file_resp = requests.get(
        f"{BALE_API_BASE}/file/bot{BALE_BOT_TOKEN}/{file_path}",
        timeout=60,
    )
    file_resp.raise_for_status()

    extension = os.path.splitext(file_path)[1] or (".jpg" if file_kind == "photo" else ".pdf")
    download_name = f"{newspaper_name}-{submission_id}{extension}"

    mimetype = "application/pdf" if extension.lower() == ".pdf" else "image/jpeg"

    return send_file(
        io.BytesIO(file_resp.content),
        as_attachment=False,
        download_name=download_name,
        mimetype=mimetype,
    )


@app.route("/submission/<int:submission_id>/delete", methods=["POST"])
@login_required
def delete_submission(submission_id):
    db = get_session()
    try:
        submission = db.query(Submission).get(submission_id)
        if submission is None:
            flash("این ارسال پیدا نشد.")
            return redirect(url_for("index"))
        db.delete(submission)
        db.commit()
        flash("رکورد حذف شد.")
    finally:
        db.close()
    return redirect(url_for("index"))


@app.route("/export.xlsx")
@login_required
def export_excel():
    import openpyxl
    from openpyxl.utils import get_column_letter

    db = get_session()
    try:
        submissions = db.query(Submission).order_by(Submission.newspaper_name).all()
    finally:
        db.close()

    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "همه اخبار"
    headers = [
        "روزنامه",
        "خبرنگار",
        "تیتر",
        "تاریخ چاپ",
        "صفحه اول",
        "امتیاز اولیه",
        "امتیاز نهایی",
        "یادداشت ادمین",
        "تاریخ ثبت",
    ]
    ws.append(headers)
    for s in submissions:
        ws.append(
            [
                s.newspaper_name,
                s.journalist_name,
                s.headline,
                s.publish_date,
                "بله" if s.is_front_page else ("خیر" if s.is_front_page is False else ""),
                s.auto_score,
                s.final_score if s.final_score is not None else "",
                s.admin_note or "",
                s.created_at.strftime("%Y-%m-%d %H:%M") if s.created_at else "",
            ]
        )
    for i, _ in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = 22

    ws2 = wb.create_sheet("جمع امتیاز روزنامه‌ها")
    ws2.append(["روزنامه", "تعداد خبر", "جمع امتیاز"])
    totals: dict[str, list[int]] = {}
    for s in submissions:
        row = totals.setdefault(s.newspaper_name, [0, 0])
        row[0] += 1
        row[1] += s.effective_score
    for name, (count, total_score) in sorted(totals.items(), key=lambda x: x[1][1], reverse=True):
        ws2.append([name, count, total_score])
    for i in range(1, 4):
        ws2.column_dimensions[get_column_letter(i)].width = 24

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"gozaresh-khabarnegaran-{datetime.date.today().isoformat()}.xlsx"
    return send_file(
        buffer,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port, debug=False) 
